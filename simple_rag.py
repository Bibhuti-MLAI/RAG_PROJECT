"""
simple_rag.py
=============
Retrieval-Augmented Generation (RAG) engine for the Legal Document Assistant.

How a question gets answered
----------------------------
1. INGEST    Every file in `legal_documents/` is read as plain text
             (PDF / TXT / CSV / DOCX). Images are first converted to a
             `<name>_extracted.txt` file with Tesseract OCR.
2. CHUNK     The text is split into ~800-character chunks that overlap by
             200 characters, so a sentence cut at a boundary is not lost.
3. EMBED     Each chunk is turned into a vector by a small local model
             (sentence-transformers/all-MiniLM-L6-v2) and stored in a Chroma
             vector database on disk (`legal_vectordb/`).
4. RETRIEVE  The question is embedded the same way and the most similar
             chunks are fetched (a few per document, so answers can cite
             several documents).
5. GENERATE  The chunks, the recent chat history and the question are sent to
             an LLM on Together AI, which is instructed to answer ONLY from
             those chunks and to name the document each fact came from.

Design notes
------------
* Every chunk gets a deterministic id `<file name>::<chunk number>`, and a
  file's old chunks are deleted before it is re-indexed. This guarantees a
  document is never stored twice in the vector database.
* Uploading a document indexes only that document (no full rebuild).
* On start-up `sync_index()` compares the folder with the database, so files
  added or removed while the server was stopped are picked up automatically.
* Follow-up questions ("what about its payment terms?") are rewritten into a
  stand-alone question using the chat history before searching.
"""

import csv
import datetime
import logging
import os
import threading
from typing import Dict, List, Optional, Tuple

# Chroma sends anonymous usage telemetry by default; switch it off.
# This must be set before chromadb is imported (langchain_chroma imports it).
os.environ.setdefault("ANONYMIZED_TELEMETRY", "False")

from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Settings
# ---------------------------------------------------------------------------

DEFAULT_EMBEDDINGS_MODEL = "sentence-transformers/all-MiniLM-L6-v2"
# The original "...-Turbo-Free" variant has been retired by Together AI
# (it now returns "model_not_available"); this is the same model, paid per token.
DEFAULT_LLM_MODEL = "meta-llama/Llama-3.3-70B-Instruct-Turbo"

CHUNK_SIZE = 800          # characters per chunk
CHUNK_OVERLAP = 200       # characters shared by neighbouring chunks
FETCH_K = 20              # chunks pulled from the vector DB per question
TOP_K = 8                 # chunks actually sent to the LLM
MAX_CHUNKS_PER_DOC = 3    # cap per document, so one big PDF can't crowd out the rest
HISTORY_TURNS = 5         # past question/answer pairs remembered per session

DOCUMENT_EXTENSIONS = [".pdf", ".txt", ".csv", ".docx"]
IMAGE_EXTENSIONS = [".png", ".jpg", ".jpeg", ".tiff", ".tif", ".bmp", ".gif"]

NOT_ENOUGH_INFO = (
    "I don't have enough information in the provided documents to answer "
    "this question. Please upload relevant documents or try a different question."
)


# ---------------------------------------------------------------------------
# Prompts sent to the LLM
# ---------------------------------------------------------------------------

SYSTEM_PROMPT = (
    "You are a legal document assistant. You answer questions strictly from "
    "the document excerpts you are given and you never use outside knowledge."
)

ANSWER_PROMPT = """Conversation so far (use it only to understand follow-up questions such as "it" or "its"):
{history}

Document excerpts:
{context}

Question: {question}

Instructions:
1. Use ONLY the document excerpts above. If they do not contain the answer, reply exactly:
   "{not_enough_info}"
2. Otherwise start with: "I've reviewed the documents and found relevant information."
3. For each document that actually contains relevant information, write a section that starts with
   "According to <document name>:" using the name shown in [Document: ...].
   Do not mention documents that have nothing relevant.
4. For a follow-up question, stay with the company or document discussed earlier unless the
   question clearly asks about something else.
5. End with a short "Summary:" of the key points."""

# Used to turn a follow-up ("what about its termination clause?") into a
# question that can be searched on its own ("What is the termination clause
# of the PureWell Labs agreement?").
CONDENSE_PROMPT = """Rewrite the follow-up question so it can be understood without the conversation.
Rules:
- Only replace vague references such as "it", "its", "they", "this" or "that agreement" with the
  company or document they refer to in the conversation.
- If the question already names its own company, document or topic, return it EXACTLY unchanged.
- Never add details, explanations or other documents that the question did not ask about.
- Keep it as short as the original question.
Return ONLY the question, nothing else.

Conversation:
{history}

Follow-up question: {question}"""


# ---------------------------------------------------------------------------
# OCR (image -> text)
# ---------------------------------------------------------------------------

def detect_ocr() -> Tuple[bool, str]:
    """Check whether image OCR can be used.

    OCR needs two things: the `pytesseract` Python package and the Tesseract
    program itself. Returns (available, human-readable message).
    """
    try:
        import pytesseract  # noqa: F401
        from PIL import Image  # noqa: F401
    except ImportError:
        return False, "OCR disabled: install the 'pytesseract' and 'Pillow' Python packages"

    import pytesseract

    # Let the user point at tesseract.exe explicitly; otherwise try the
    # default Windows install location (on Linux it is normally on PATH).
    candidates = [os.environ.get("TESSERACT_CMD", "")]
    if os.name == "nt":
        candidates.append(r"C:\Program Files\Tesseract-OCR\tesseract.exe")
    for path in candidates:
        if path and os.path.exists(path):
            pytesseract.pytesseract.tesseract_cmd = path
            break

    try:
        version = pytesseract.get_tesseract_version()
        return True, f"OCR available (Tesseract {version})"
    except Exception:
        return False, "OCR disabled: the Tesseract program was not found (install it or set TESSERACT_CMD)"


def ocr_image_to_text_file(image_path: str, output_dir: str) -> str:
    """Read the text in an image and save it as `<image name>_extracted.txt`.

    Returns the path of the new text file. Raises an exception if OCR fails
    or finds no text, and in that case the original image is left untouched.
    The image is deleted only after the text file has been written.
    """
    import pytesseract
    from PIL import Image

    with Image.open(image_path) as img:
        # Put transparent images on a white background, then convert to
        # grayscale, which is what Tesseract works best with.
        if img.mode in ("RGBA", "LA", "P"):
            rgba = img.convert("RGBA")
            background = Image.new("RGBA", rgba.size, "white")
            img = Image.alpha_composite(background, rgba)
        raw_text = pytesseract.image_to_string(img.convert("L"))

    # Drop empty lines and lines that are only symbols (typical OCR noise).
    lines = [line.strip() for line in raw_text.splitlines()]
    lines = [line for line in lines if len(line) > 1 and any(ch.isalnum() for ch in line)]
    if not lines:
        raise ValueError("No text could be read from the image")

    image_name = os.path.basename(image_path)
    output_path = os.path.join(output_dir, f"{os.path.splitext(image_name)[0]}_extracted.txt")
    header = (
        f"# Text extracted from image: {image_name}\n"
        f"# Extraction date: {datetime.datetime.now():%Y-%m-%d %H:%M:%S}\n\n"
    )
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(header + "\n".join(lines))

    os.remove(image_path)
    logger.info("OCR: %s -> %s", image_name, os.path.basename(output_path))
    return output_path


# ---------------------------------------------------------------------------
# File readers: each one returns the whole file as a single string
# ---------------------------------------------------------------------------

def read_txt(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def read_pdf(path: str) -> str:
    import pypdf

    reader = pypdf.PdfReader(path)
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def read_docx(path: str) -> str:
    import docx  # the python-docx package

    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    # Contracts often keep terms in tables, so read table cells as well.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def read_csv(path: str) -> str:
    # Each row becomes a block of "column: value" lines.
    with open(path, "r", encoding="utf-8", errors="replace", newline="") as f:
        rows = list(csv.DictReader(f))
    return "\n\n".join(
        "\n".join(f"{key}: {value}" for key, value in row.items()) for row in rows
    )


FILE_READERS = {".txt": read_txt, ".pdf": read_pdf, ".docx": read_docx, ".csv": read_csv}


# ---------------------------------------------------------------------------
# The RAG engine
# ---------------------------------------------------------------------------

class SimpleLegalRAG:
    """Indexes the documents folder and answers questions about it."""

    def __init__(
        self,
        docs_dir: str = "legal_documents",
        persist_dir: str = "legal_vectordb",
        embeddings_model: str = DEFAULT_EMBEDDINGS_MODEL,
        together_model: Optional[str] = None,
        together_api_key: Optional[str] = None,
        collection_name: str = "legal_documents",
    ):
        self.docs_dir = docs_dir
        self.persist_dir = persist_dir
        self.embeddings_model = embeddings_model
        self.together_model = together_model or os.environ.get("TOGETHER_MODEL", DEFAULT_LLM_MODEL)
        self.together_api_key = together_api_key or os.environ.get("TOGETHER_API_KEY")
        self.collection_name = collection_name

        self.supported_extensions = list(DOCUMENT_EXTENSIONS)
        self.image_extensions = list(IMAGE_EXTENSIONS)
        self.ocr_available, self.ocr_message = detect_ocr()

        # Created in setup()
        self.embeddings = None
        self.vectordb: Optional[Chroma] = None
        self.llm_client = None

        # session id -> list of (question, answer), oldest first
        self.chat_histories: Dict[str, List[Tuple[str, str]]] = {}

        # FastAPI runs requests in parallel threads; these locks stop two
        # uploads from changing the index at the same time and keep the
        # chat history dictionary consistent.
        self._index_lock = threading.RLock()
        self._history_lock = threading.Lock()

        self.splitter = RecursiveCharacterTextSplitter(chunk_size=CHUNK_SIZE, chunk_overlap=CHUNK_OVERLAP)

        os.makedirs(self.docs_dir, exist_ok=True)
        os.makedirs(self.persist_dir, exist_ok=True)

    # ----- start-up -------------------------------------------------------

    def setup(self) -> None:
        """Load the embedding model, open the vector DB, connect the LLM and
        bring the index in line with the documents folder."""
        logger.info("Loading embedding model %s", self.embeddings_model)
        self.embeddings = HuggingFaceEmbeddings(model_name=self.embeddings_model)

        self.vectordb = Chroma(
            collection_name=self.collection_name,
            embedding_function=self.embeddings,
            persist_directory=self.persist_dir,
            collection_metadata={"hnsw:space": "cosine"},
        )

        if self.together_api_key:
            from together import Together

            self.llm_client = Together(api_key=self.together_api_key)
            logger.info("LLM: %s via Together AI", self.together_model)
        else:
            logger.warning("TOGETHER_API_KEY is not set - questions cannot be answered until it is")

        logger.info(self.ocr_message)
        summary = self.sync_index()
        logger.info("Index ready: %s", summary)

    # ----- documents on disk ---------------------------------------------

    def get_supported_file_types(self) -> List[str]:
        """Extensions accepted for upload (images only when OCR works)."""
        if self.ocr_available:
            return self.supported_extensions + self.image_extensions
        return list(self.supported_extensions)

    def is_image_file(self, file_name: str) -> bool:
        return os.path.splitext(file_name)[1].lower() in self.image_extensions

    def is_ocr_available(self) -> bool:
        return self.ocr_available

    def list_document_files(self) -> List[str]:
        """Names of the indexable documents in the documents folder."""
        names = []
        for name in os.listdir(self.docs_dir):
            path = os.path.join(self.docs_dir, name)
            if not os.path.isfile(path) or name.startswith("."):
                continue  # skip folders, hidden files and in-progress uploads
            if os.path.splitext(name)[1].lower() in self.supported_extensions:
                names.append(name)
        return sorted(names)

    def load_file_text(self, file_name: str) -> str:
        """Read one document from the documents folder as plain text."""
        extension = os.path.splitext(file_name)[1].lower()
        reader = FILE_READERS.get(extension)
        if reader is None:
            raise ValueError(f"Unsupported file type: {extension}")
        return reader(os.path.join(self.docs_dir, file_name))

    # ----- vector index ---------------------------------------------------

    def _file_version(self, file_name: str) -> str:
        """The file's modification time in nanoseconds, as text.

        Stored with every chunk to detect changed files. It is kept as a
        string because Chroma rounds float metadata, which made unchanged
        files look modified.
        """
        return str(os.stat(os.path.join(self.docs_dir, file_name)).st_mtime_ns)

    def _indexed_files(self) -> Dict[str, Optional[str]]:
        """{file name: file version it had when indexed} for every file in the DB."""
        metadatas = self.vectordb.get(include=["metadatas"])["metadatas"]
        return {m["file_name"]: m.get("file_version") for m in metadatas if m and "file_name" in m}

    def index_file(self, file_name: str) -> int:
        """(Re)index a single document and return how many chunks were stored.

        Any chunks already stored for this file are removed first, so a file
        is never duplicated in the database.
        """
        text = self.load_file_text(file_name)
        if not text.strip():
            raise ValueError(f"No readable text found in {file_name}")

        document_name = os.path.splitext(file_name)[0]
        version = self._file_version(file_name)
        doc_type = "txt_from_image" if file_name.endswith("_extracted.txt") else os.path.splitext(file_name)[1][1:].lower()

        chunks = []
        for i, piece in enumerate(self.splitter.split_text(text)):
            chunks.append(Document(
                # The document name is written into the chunk text too, so the
                # LLM always knows which document a passage came from.
                page_content=f"[From: {document_name}]\n{piece}",
                metadata={
                    "file_name": file_name,
                    "document_name": document_name,
                    "document_type": doc_type,
                    "chunk_index": i,
                    "file_version": version,
                },
            ))
        ids = [f"{file_name}::{i}" for i in range(len(chunks))]

        with self._index_lock:
            self.remove_from_index(file_name)
            self.vectordb.add_documents(chunks, ids=ids)

        logger.info("Indexed %s (%d chunks)", file_name, len(chunks))
        return len(chunks)

    def remove_from_index(self, file_name: str) -> int:
        """Delete all chunks of one file from the vector DB (exact name match)."""
        with self._index_lock:
            ids = self.vectordb.get(where={"file_name": file_name})["ids"]
            if ids:
                self.vectordb.delete(ids=ids)
        return len(ids)

    def document_exists_in_vectordb(self, file_name: str) -> bool:
        return bool(self.vectordb.get(where={"file_name": file_name}, limit=1)["ids"])

    def sync_index(self) -> Dict[str, List[str]]:
        """Make the vector DB match the documents folder.

        * new or modified files are indexed
        * files that no longer exist have their chunks removed
        """
        with self._index_lock:
            on_disk = {name: self._file_version(name) for name in self.list_document_files()}
            indexed = self._indexed_files()

            added, removed, failed = [], [], []
            for name, version in on_disk.items():
                if indexed.get(name) != version:
                    try:
                        self.index_file(name)
                        added.append(name)
                    except Exception as e:
                        logger.error("Could not index %s: %s", name, e)
                        failed.append(name)
            for name in indexed:
                if name not in on_disk:
                    self.remove_from_index(name)
                    removed.append(name)

        return {"added": added, "removed": removed, "failed": failed}

    def rebuild_index(self) -> Dict[str, List[str]]:
        """Empty the vector DB and index every document again from scratch."""
        with self._index_lock:
            self.vectordb.reset_collection()
            return self.sync_index()

    # ----- adding / removing documents -----------------------------------

    def add_document(self, file_path: str) -> Dict:
        """Index a file that was just saved into the documents folder.

        Images are converted to `<name>_extracted.txt` with OCR first.
        Raises an exception if the file cannot be processed.
        """
        file_name = os.path.basename(file_path)
        converted_from_image = False

        if self.is_image_file(file_name):
            if not self.ocr_available:
                raise RuntimeError(self.ocr_message)
            file_name = os.path.basename(ocr_image_to_text_file(file_path, self.docs_dir))
            converted_from_image = True

        chunk_count = self.index_file(file_name)
        message = f"Document processed successfully ({chunk_count} chunks indexed)"
        if converted_from_image:
            message = f"Image converted to text with OCR and indexed ({chunk_count} chunks)"

        return {
            "success": True,
            "message": message,
            "processed_file": file_name,
            "converted_from_image": converted_from_image,
            "chunks": chunk_count,
        }

    def remove_document(self, file_name: str) -> int:
        """Delete a document from the index and from disk. Returns the number of chunks removed."""
        removed = self.remove_from_index(file_name)
        os.remove(os.path.join(self.docs_dir, file_name))
        logger.info("Deleted %s (%d chunks)", file_name, removed)
        return removed

    # ----- retrieval ------------------------------------------------------

    def retrieve(self, query: str) -> List[Document]:
        """Return the chunks most relevant to `query`.

        Fetches FETCH_K candidates by similarity, drops exact duplicates and
        keeps at most MAX_CHUNKS_PER_DOC per document so answers can draw on
        several documents. The best matches come first.
        """
        candidates = self.vectordb.similarity_search(query, k=FETCH_K)

        selected, seen_text, per_doc = [], set(), {}
        for doc in candidates:
            name = doc.metadata.get("document_name", "Unknown")
            if doc.page_content in seen_text or per_doc.get(name, 0) >= MAX_CHUNKS_PER_DOC:
                continue
            seen_text.add(doc.page_content)
            per_doc[name] = per_doc.get(name, 0) + 1
            selected.append(doc)
            if len(selected) == TOP_K:
                break
        return selected

    # ----- LLM ------------------------------------------------------------

    def _chat(self, user_prompt: str, max_tokens: int = 2048) -> str:
        """Send one prompt to the LLM and return its reply text."""
        response = self.llm_client.chat.completions.create(
            model=self.together_model,
            messages=[
                {"role": "system", "content": SYSTEM_PROMPT},
                {"role": "user", "content": user_prompt},
            ],
            temperature=0.0,  # deterministic, no creativity: we want facts from the documents
            max_tokens=max_tokens,
        )
        return (response.choices[0].message.content or "").strip()

    @staticmethod
    def _format_history(history: List[Tuple[str, str]]) -> str:
        if not history:
            return "(no previous messages)"
        return "\n".join(f"User: {q}\nAssistant: {a}" for q, a in history)

    def _standalone_question(self, question: str, history: List[Tuple[str, str]]) -> str:
        """Rewrite a follow-up question so it can be searched without the chat history."""
        if not history:
            return question
        try:
            prompt = CONDENSE_PROMPT.format(history=self._format_history(history), question=question)
            rewritten = self._chat(prompt, max_tokens=200)
            return rewritten or question
        except Exception as e:
            # Fall back to searching with the previous question added for context.
            logger.warning("Could not rewrite follow-up question: %s", e)
            return f"{history[-1][0]} {question}"

    # ----- answering questions -------------------------------------------

    def query(self, question: str, session_id: str = "default") -> Dict:
        """Answer a question. Returns {"question", "answer", "sources", "search_query"}.

        Raises an exception if the LLM call fails, so the API can report it.
        """
        result = {"question": question, "answer": "", "sources": [], "search_query": question}

        if self.llm_client is None:
            result["answer"] = "The language model is not configured. Set TOGETHER_API_KEY and restart the server."
            return result
        if not self.list_document_files():
            result["answer"] = "There are no documents in the system yet. Please upload legal documents to get answers to your questions."
            return result

        history = self.get_history(session_id)

        # 1. Turn follow-ups into a self-contained question, then search with it.
        search_query = self._standalone_question(question, history)
        result["search_query"] = search_query
        docs = self.retrieve(search_query)
        if not docs:
            result["answer"] = NOT_ENOUGH_INFO
            return result

        # 2. Group the chunks by document and build the context block.
        by_document: Dict[str, List[Document]] = {}
        for doc in docs:
            by_document.setdefault(doc.metadata.get("document_name", "Unknown"), []).append(doc)
        context = "\n\n====================\n\n".join(
            f"[Document: {name}]\n" + "\n\n".join(d.page_content for d in chunks)
            for name, chunks in by_document.items()
        )

        # 3. Ask the LLM.
        prompt = ANSWER_PROMPT.format(
            history=self._format_history(history),
            context=context,
            question=question,
            not_enough_info=NOT_ENOUGH_INFO,
        )
        answer = self._chat(prompt)
        self._add_to_history(session_id, question, answer)

        # 4. One source entry per document, showing a preview of its best chunk.
        for name, chunks in by_document.items():
            preview = chunks[0].page_content.split("\n", 1)[-1]  # drop the "[From: ...]" line
            result["sources"].append({
                "source": name,
                "content": preview[:200] + ("..." if len(preview) > 200 else ""),
            })

        result["answer"] = answer
        return result

    # ----- conversation history -----------------------------------------

    def get_history(self, session_id: str) -> List[Tuple[str, str]]:
        with self._history_lock:
            return list(self.chat_histories.get(session_id, []))

    def _add_to_history(self, session_id: str, question: str, answer: str) -> None:
        with self._history_lock:
            turns = self.chat_histories.setdefault(session_id, [])
            turns.append((question, answer))
            del turns[:-HISTORY_TURNS]  # keep only the most recent turns

    def session_message_count(self, session_id: str) -> int:
        with self._history_lock:
            return 2 * len(self.chat_histories.get(session_id, []))  # a question + an answer per turn

    def clear_session(self, session_id: str) -> bool:
        with self._history_lock:
            return self.chat_histories.pop(session_id, None) is not None

    # ----- maintenance / status -----------------------------------------

    def reset(self) -> Dict[str, List[str]]:
        """Forget all conversations and rebuild the index from the documents
        folder. Documents on disk are kept."""
        with self._history_lock:
            self.chat_histories.clear()
        return self.rebuild_index()

    def vector_count(self) -> int:
        return len(self.vectordb.get(include=[])["ids"]) if self.vectordb else 0

    def get_system_status(self) -> Dict:
        files = self.list_document_files()
        return {
            "total_documents": len(files),
            "text_from_images": sum(1 for f in files if f.endswith("_extracted.txt")),
            "indexed_documents": len(self._indexed_files()) if self.vectordb else 0,
            "vector_entries": self.vector_count(),
            "ocr_available": self.ocr_available,
            "ocr_message": self.ocr_message,
            "supported_extensions": self.get_supported_file_types(),
            "active_sessions": len(self.chat_histories),
            "llm_model": self.together_model,
            "system_components": {
                "embeddings_initialized": self.embeddings is not None,
                "vectordb_initialized": self.vectordb is not None,
                "llm_initialized": self.llm_client is not None,
            },
        }
